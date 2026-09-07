from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_style(doc):
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

def create_unit_1():
    doc = Document()
    set_style(doc)
    doc.add_heading('Unit 1: Executive Overview & Value Proposition', 0)
    
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph("ENVALIOR EXECUTIVE SUMMARY CONFIDENTIAL")
    doc.add_paragraph(
        "The Touchless Order Creation solution automates the conversion of inbound Purchase Orders (POs) into SAP Sales Orders (SOs), "
        "eliminating manual data entry. The process follows a straight-through pipeline: a customer sends a PO by email → "
        "the AI engine extracts and validates the data → Celonis applies business rules → a Sales Order is automatically created in SAP. "
        "If anything fails, the PO is routed to the regional Customer Service team for manual handling."
    )
    
    doc.add_heading('2. Core Functionality', level=1)
    functions = [
        "Monitors central sales inbox (salesorders@envalior.com) for incoming PO emails.",
        "Extracts key data from PDF attachments using AI (customer, PO#, materials, quantities, dates).",
        "Matches extracted data against SAP Master Data (KNA1, KNVP, KNMT) to resolve IDs.",
        "Applies business rules (Sales Org, Order Type, SO Splitting) using the Decision Tree & Transactional Model.",
        "Triggers SO creation in SAP EPP via Celonis Action Flows.",
        "Archives original PO documents on the Sales Order in SAP for compliance."
    ]
    for func in functions:
        doc.add_paragraph(func, style='List Bullet')

    doc.add_heading('3. Business Prerequisites', level=1)
    doc.add_paragraph(
        "• Azure App Registration: Mail.ReadWrite and Mail.Send permissions.\n"
        "• Celonis Connection: Active Data Pool with write access (PO_EXTRACTION_RESULTS).\n"
        "• SAP Master Data: Customer and Material Info Records synced to Celonis.\n"
        "• Middleware: Webmethods/BAPI configured for PO archiving."
    )

    doc.save('Unit_1_Executive_Overview_PROPER.docx')

def create_unit_2():
    doc = Document()
    set_style(doc)
    doc.add_heading('Unit 2: Technical Architecture & Infrastructure', 0)
    
    doc.add_heading('1. System Architecture Overview', level=1)
    doc.add_paragraph(
        "The system operates through a fully automated multi-stage pipeline using MS Graph API, Azure cloud, "
        "PaddleOCR, Claude 4 Sonnet, and Celonis Action Flows."
    )
    doc.add_paragraph("[PLACEMENT FOR ARCHITECTURE DIAGRAM]")

    doc.add_heading('2. Detailed Pipeline Stages', level=1)
    stages = [
        ("Stage 1: Inbox Monitoring", "Outlook Inbox polled every 60 seconds for unread emails with PDF attachments."),
        ("Stage 2: Outlook Poller", "Retrieves attachments via Graph API, pushes to po-processing-queue and Blob Storage."),
        ("Stage 3: Azure Cloud Storage", "Azure Storage Queue triggers downstream processing; Blob Storage (input-po) stores raw files."),
        ("Stage 4: PaddleOCR", "Converts scanned/digital PDFs into raw text output."),
        ("Stage 5: Claude 4 Sonnet (LLM)", "Structured JSON extraction (Customer, PO#, Line Items, Dates)."),
        ("Stage 6: Knowledge Base", "SQLite enrichment (knowledge_base.db) using SAP master data (KNA1, KNVP, KNMT)."),
        ("Stage 7: Fuzzy Match Engine", "Resolves Sold-To/Ship-To IDs via TheFuzz algorithms (token_sort_ratio)."),
        ("Stage 8: Sales Order Builder", "Applies Sales Org/Order Type and Grouping Rules (Rule X / Rule O)."),
        ("Stage 9: SAP Ingestion", "Final JSON payload picked up by Celonis for SAP EPP writeback.")
    ]
    for stage_h, stage_d in stages:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f"{stage_h}: ").bold = True
        p.add_run(stage_d)

    doc.save('Unit_2_Technical_Architecture_PROPER.docx')

def create_unit_3():
    doc = Document()
    set_style(doc)
    doc.add_heading('Unit 3: Functional Design & Decision Intelligence', 0)
    
    doc.add_heading('1. Order Type Identification (Decision Tree)', level=1)
    doc.add_paragraph(
        "Order type is determined based on the following hierarchical logic:"
    )
    logic = [
        "1. Decision Node: Does the customer have a consignment block? -> Yes: Order Type = TA.",
        "2. Decision Node: Is 'consignment' mentioned in the email? -> Yes: Check History for KB/ZKB.",
        "3. Fallback: Full historical check for Customer/Material combo from transaction history record."
    ]
    for l in logic:
        doc.add_paragraph(l, style='List Number')

    doc.add_heading('2. Customer & Material Resolution', level=1)
    doc.add_paragraph(
        "• Sold-To: Matches filename metadata or fuzzy-matched extracted name against KNA1.\n"
        "• Ship-To: Fuzzy-matched extracted address against linked KNVP records.\n"
        "• Material Mapping: 3-tier fallback (Case-insensitive -> Startswith -> Reverse Prefix Matching)."
    )
    
    doc.add_heading('3. Mapping Maintenance Rules', level=1)
    doc.add_paragraph(
        "Rules are managed via the Mapping Excel flat file containing columns: Sales Org, Customer, Material, "
        "Customer Material Number, Order Type, and One SO per PO (X/O)."
    )

    doc.save('Unit_3_Functional_Logic_PROPER.docx')

def create_unit_4():
    doc = Document()
    set_style(doc)
    doc.add_heading('Unit 4: Operational Management & Exception Routing', 0)
    
    doc.add_heading('1. Exception Handling Strategy', level=1)
    doc.add_paragraph(
        "Unprocessable records (review_required: true) follow a defined regional handoff:"
    )
    doc.add_paragraph(
        "• AI Failure: Low confidence POs sent to Regional CS Inbox.\n"
        "• Celonis Failure: SAP-push errors (credit blocks, validation) reported back to CS.\n"
        "• Feedback Loop: Manual corrections improve terminal classification logic and update CUSTOMER_ALIASES."
    )
    
    doc.add_heading('2. Maintenance Guide', level=1)
    steps = [
        "Updating Celonis Master Data: Daily sync from SAP.",
        "Managing Excel Mapping: Adding new Sold-To entries and assigning Sales Org/Order Type.",
        "Material Cross-References: Mapping customer part numbers (e.g., V123) to SAP IDs (e.g., MAT-9000)."
    ]
    for s in steps:
        doc.add_paragraph(s, style='List Bullet')

    doc.add_heading('3. Support Model', level=1)
    doc.add_paragraph("Level 1 support handles inbox oversight, Level 2 handles mapping/alias updates, Level 3 handles engine calibration.")

    doc.save('Unit_4_Support_Maintenance_PROPER.docx')

def create_unit_5():
    doc = Document()
    set_style(doc)
    doc.add_heading('Unit 5: Validation & Quality Suite', 0)
    
    doc.add_heading('1. Test Methodology', level=1)
    doc.add_paragraph("Functional and Technical validation ensures system reliability.")
    
    doc.add_heading('2. Functional Unit Testing (FUT)', level=1)
    doc.add_paragraph(
        "Validates data logic: Customer resolution, material mapping fallbacks, and SO grouping (Rule X/O)."
    )
    
    doc.add_heading('3. Technical Unit Testing (TUT)', level=1)
    doc.add_paragraph(
        "Verifies backend script execution: Graph API connectivity, Blob Storage retrieval, and Celonis pycelonis SDK push."
    )

    doc.add_heading('4. User Acceptance Scenarios', level=1)
    scenarios = [
        "Scenario 1: Happy Path PO -> Automated SO creation in SAP EPP.",
        "Scenario 2: Missing Mapping -> Routing to Regional CS Desk.",
        "Scenario 3: Multi-Line PO -> Splitting into separate SOs via Rule O.",
        "Scenario 4: Alias Match -> Resolving non-canonical names via CUSTOMER_ALIASES."
    ]
    for sc in scenarios:
        doc.add_paragraph(sc, style='List Bullet')

    doc.save('Unit_5_Validation_Testing_PROPER.docx')

if __name__ == "__main__":
    create_unit_1()
    create_unit_2()
    create_unit_3()
    create_unit_4()
    create_unit_5()
    print("All PROPER units created successfully.")
