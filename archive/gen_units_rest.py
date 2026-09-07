from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_unit_2():
    doc = Document()
    doc.add_heading('Unit 2: Technical Architecture & Middleware', 0)
    
    doc.add_heading('1. System Architecture', level=1)
    doc.add_paragraph("[IMAGE PLACEHOLDER: TECHNICAL_ARCHITECTURE_DIAGRAM]")
    doc.add_paragraph(
        "The architecture is built on a hybrid cloud model using Azure, Celonis, and Envalior's SAP EPP environment."
    )
    
    doc.add_heading('2. Core Infrastructure Components', level=1)
    components = {
        "Microsoft Graph API": "Bridges the gap between the customer's email and our AI engine. Handles polling, archiving, and error-routing handed off to CS teams.",
        "Azure Blob Storage": "Acts as a scalable repository for all raw PO images and PDFs, ensuring high availability and persistence.",
        "PaddleOCR & Claude LLM": "A powerful multi-stage extraction core. OCR handles the visual layer, and the LLM (Large Language Model) interprets the semantic data into structured JSON.",
        "Robona RPA": "Automates the final 'Attachment' step, ensuring the source of truth is always available within the SAP Sales Order record."
    }
    for comp, desc in components.items():
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f"{comp}: ").bold = True
        p.add_run(desc)

    doc.add_heading('3. Prerequisites for Deployment', level=1)
    prereqs = [
        "Azure App Registration with Mail.ReadWrite permissions.",
        "Active Celonis Sandbox/Production Data Pool access.",
        "SAP Master Data Sync (KNA1, KNVP, KNMT) via pycelonis.",
        "Microsoft Graph API credentials (Client ID, Secret, Tenant ID)."
    ]
    for pr in prereqs:
        doc.add_paragraph(pr, style='List Bullet')

    doc.save('Unit_2_Technical_Architecture.docx')

def create_unit_3():
    doc = Document()
    doc.add_heading('Unit 3: Functional Logic & Decision Tree', 0)
    
    doc.add_heading('1. Order Type Decision Logic', level=1)
    doc.add_paragraph(
        "To achieve a touchless outcome, the system replaces manual Excel lookups with a dynamic Decision Tree. "
        "The Order Type (TA vs. KB) is determined based on real-time PO context."
    )
    doc.add_paragraph("[IMAGE PLACEHOLDER: ORDER_TYPE_DECISION_TREE]")
    
    doc.add_heading('2. Tiered Material Mapping', level=1)
    doc.add_paragraph(
        "Material mapping ensures that the customer's material terminology matches our internal SAP IDs through 3 fallback tiers:\n"
        "1. Direct Match: Exact alphanumeric comparison against the Knowledge Base.\n"
        "2. Aliasing: Using the CS-maintained Alias dictionary for known terminology differences.\n"
        "3. Fuzzy Search: Intelligent string matching to resolve typos or minor variations."
    )

    doc.add_heading('3. SO Grouping Rules', level=1)
    doc.add_paragraph(
        "The logic supports business-defined splitting:\n"
        "- Rule X: One Sales Order created per Purchase Order.\n"
        "- Rule O: Each individual line item triggers a separate Sales Order (for multi-delivery scenarios)."
    )

    doc.save('Unit_3_Functional_Logic.docx')

def create_unit_4():
    doc = Document()
    doc.add_heading('Unit 4: Operational & CS Support Guide', 0)
    
    doc.add_heading('1. Regional CS Workflow', level=1)
    doc.add_paragraph(
        "The system acts as a digital assistant for the Customer Service team, filtering out 100% of 'Perfect POs' "
        "and presenting only complex cases for human review."
    )
    
    doc.add_heading('2. Exception Routing', level=1)
    doc.add_paragraph(
        "When AI confidence is low, the pipeline automatically routes the file:\n"
        "- EMEA Desk: Routing via Graph API to regional mailbox.\n"
        "- APAC/AMS Desks: Follow-on routing based on Sold-To region detection."
    )

    doc.add_heading('3. AI Feedback Loop', level=1)
    doc.add_paragraph(
        "When a manual correction is made by CS, the feedback is fed back into the 'Enrichment' layer. "
        "This ensures that next time a similar PO arrives, the AI can handle it touchlessly."
    )

    doc.save('Unit_4_Support_Maintenance.docx')

def create_unit_5():
    doc = Document()
    doc.add_heading('Unit 5: Verification & Quality Assurance', 0)
    
    doc.add_heading('1. UAT Success Criteria', level=1)
    doc.add_paragraph(
        "This unit summarizes the successful testing cycles conducted to validate the Touchless system."
    )
    
    doc.add_heading('2. Functional Unit Testing (FUT)', level=1)
    doc.add_paragraph(
        "Validates that data extraction matches the PO source with high precision. "
        "Results: Passed 98% Extraction accuracy during internal sprints."
    )

    doc.add_heading('3. User Acceptance Testing (UAT)', level=1)
    doc.add_paragraph(
        "Final validation by business users across multiple regions. "
        "Scenarios included single-line files, multi-page PDFs, and multilingual inputs."
    )

    doc.save('Unit_5_Validation_Testing.docx')

if __name__ == "__main__":
    create_unit_2()
    create_unit_3()
    create_unit_4()
    create_unit_5()
    print("Units 2-5 created.")
