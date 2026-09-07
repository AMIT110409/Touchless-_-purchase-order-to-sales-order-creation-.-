"""
create_uat_functional.py
Generates: Touchless_Order_Creation_UAT_Functional.docx
Business-friendly UAT - no technical commands, no API names, no code.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DARK_BLUE  = RGBColor(0x00, 0x35, 0x6B)
MID_BLUE   = RGBColor(0x00, 0x6D, 0xC6)
LIGHT_BLUE = RGBColor(0xD6, 0xE8, 0xF7)
GREEN      = RGBColor(0xD6, 0xF0, 0xD6)
RED        = RGBColor(0xFC, 0xE4, 0xE4)
AMBER      = RGBColor(0xFF, 0xF2, 0xCC)
GREY       = RGBColor(0xF2, 0xF2, 0xF2)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GREY  = RGBColor(0x40, 0x40, 0x40)

def shd(cell, rgb):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    s    = OxmlElement("w:shd")
    s.set(qn("w:val"),   "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"),  f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}")
    tcPr.append(s)

def col_widths(table, widths):
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                cell.width = Cm(widths[i])

def hdr(table, cols, bg=DARK_BLUE, fg=WHITE, sz=9):
    row = table.rows[0]
    for i, h in enumerate(cols):
        if i >= len(row.cells): break
        c = row.cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True; r.font.size = Pt(sz); r.font.color.rgb = fg
        shd(c, bg)

def cell_text(cell, text, size=9, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color: r.font.color.rgb = color

def h1(doc, text):
    p = doc.add_heading(text, level=1)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:
        r.font.color.rgb = DARK_BLUE; r.bold = True
    return p

def h2(doc, text):
    p = doc.add_heading(text, level=2)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:
        r.font.color.rgb = MID_BLUE; r.bold = True
    return p

def para(doc, text="", size=10, bold=False, italic=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    return p

def bullet(doc, text, size=10):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text).font.size = Pt(size)

def scenario_table(doc, steps_data, row_bg=None):
    """steps_data = list of (step_no, what_to_do, what_to_check, pass_fail)"""
    t = doc.add_table(rows=len(steps_data) + 1, cols=4)
    t.style = "Table Grid"
    hdr(t, ["Step", "What to Do", "What to Check / Expect", "Pass / Fail"], DARK_BLUE, WHITE, 9)
    col_widths(t, [1.0, 5.5, 7.5, 2.0])
    for i, (num, do, check, _) in enumerate(steps_data):
        row = t.rows[i + 1]
        bg  = row_bg[i] if row_bg else (GREY if i % 2 == 0 else WHITE)
        cell_text(row.cells[0], num, 9, bold=True)
        cell_text(row.cells[1], do, 9)
        cell_text(row.cells[2], check, 9)
        cell_text(row.cells[3], "☐  Pass\n☐  Fail\n\n_______\nSigned:", 9)
        for c in row.cells:
            shd(c, bg)
    return t

def result_row(doc):
    t = doc.add_table(rows=2, cols=3)
    t.style = "Table Grid"
    hdr(t, ["Scenario Result", "Tester", "Client Sign-Off"], DARK_BLUE, WHITE, 9)
    col_widths(t, [5.0, 5.0, 6.0])
    row = t.rows[1]
    cell_text(row.cells[0], "☐  PASS — All checks confirmed\n☐  FAIL — Gap identified (record below)", 9)
    cell_text(row.cells[1], "Name: ____________________\nDate:  ____________________", 9)
    cell_text(row.cells[2], "☐  ACCEPTED\n☐  REJECTED\n\nName: ____________________\nDate:  ____________________", 9)
    for c in row.cells:
        shd(c, GREY)
    return t

# ─────────────────────────────────────────────────────────────────
def build():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin    = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin   = Cm(2.5)
        sec.right_margin  = Cm(2.5)

    # ══ COVER ═══════════════════════════════════════════════════════
    doc.add_paragraph(); doc.add_paragraph()
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tp.add_run("TOUCHLESS ORDER CREATION")
    r.bold = True; r.font.size = Pt(22); r.font.color.rgb = DARK_BLUE

    sp = doc.add_paragraph()
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sp.add_run("Functional User Acceptance Testing (UAT)")
    r2.bold = True; r2.font.size = Pt(16); r2.font.color.rgb = MID_BLUE

    doc.add_paragraph()
    sp2 = doc.add_paragraph()
    sp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = sp2.add_run("Business-Focused Test Pack — No Technical Knowledge Required")
    r3.italic = True; r3.font.size = Pt(11); r3.font.color.rgb = DARK_GREY

    doc.add_paragraph()
    mt = doc.add_table(rows=5, cols=2)
    mt.style = "Table Grid"
    mt.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta = [
        ("Document Type",   "Functional UAT — Business Sign-Off"),
        ("Project",         "Touchless Order Creation"),
        ("Prepared For",    "Client Business Team"),
        ("Version",         "2.0  |  Simplified Functional Release"),
        ("Date",            "July 2026"),
    ]
    for i, (k, v) in enumerate(meta):
        row = mt.rows[i]
        row.cells[0].text = k; row.cells[1].text = v
        shd(row.cells[0], LIGHT_BLUE)
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.bold = (c == row.cells[0])
                    r.font.color.rgb = DARK_BLUE if c == row.cells[0] else DARK_GREY
    col_widths(mt, [5.5, 10.5])
    doc.add_page_break()

    # ══ 1. HOW TO USE THIS DOCUMENT ══════════════════════════════════
    h1(doc, "1.  How to Use This Document")
    para(doc, (
        "This document is for the business team — Customer Service, Operations, and Project Leadership. "
        "You do not need any technical knowledge to complete this UAT. "
        "Your job is to observe what the system does and confirm it matches what you expect."
    ), 10)
    para(doc, (
        "The AI team will operate the system. Your team observes, checks the results listed in each "
        "step, and marks Pass or Fail. If something does not look right, record it in the Gap Log "
        "at the end of this document."
    ), 10)
    doc.add_paragraph()

    para(doc, "What the system should do — in plain terms:", 10, bold=True)
    bullet(doc, "A customer sends a Purchase Order (PO) as a PDF by email to the shared mailbox.")
    bullet(doc, "The system reads the email and PDF automatically — no one needs to open it manually.")
    bullet(doc, "It finds the customer details and product codes from the PDF.")
    bullet(doc, "If everything is recognised, it creates a Sales Order in SAP automatically.")
    bullet(doc, "If something is missing or unrecognised, it sends an alert email to the Customer Service team.")
    bullet(doc, "The original PO email is moved to the correct Outlook folder (Processed or Exception).")
    doc.add_page_break()

    # ══ 2. BEFORE YOU START — PRE-TEST CHECKLIST ═════════════════════
    h1(doc, "2.  Before You Start — Pre-Test Checklist")
    para(doc, "Tick each item before running any test scenario. If any item is not ready, UAT cannot begin.", 10)
    doc.add_paragraph()

    pc_tbl = doc.add_table(rows=8, cols=3)
    pc_tbl.style = "Table Grid"
    hdr(pc_tbl, ["#", "What to Check", "Ready?"], DARK_BLUE, WHITE, 9)
    col_widths(pc_tbl, [1.0, 12.0, 3.0])
    pcs = [
        ("PC-01", "You can open and view the shared PO mailbox in Outlook."),
        ("PC-02", "The Outlook folder list includes: 'Unprocessed POs', 'Processed POs', 'Exception POs'."),
        ("PC-03", "You have at least 3 sample PO PDF files ready to use during testing:\n"
                  "   - Test PO 1: A valid, complete PO from a known customer with a known product code.\n"
                  "   - Test PO 2: A PO with a customer name that does NOT exist in SAP.\n"
                  "   - Test PO 3: A PO from a known customer but with a product code not in SAP."),
        ("PC-04", "You can view Sales Orders in SAP (to verify that SOs are created correctly)."),
        ("PC-05", "You can access your Customer Service team email inbox (to check exception notifications)."),
        ("PC-06", "The Robona RPA service mailbox is accessible (to verify audit notifications)."),
        ("PC-07", "The AI team has confirmed the system is connected and ready to run."),
    ]
    for i, (ref, text) in enumerate(pcs):
        row = pc_tbl.rows[i + 1]
        cell_text(row.cells[0], ref, 9, bold=True)
        cell_text(row.cells[1], text, 9)
        cell_text(row.cells[2], "☐  Yes\n☐  No", 9)
        shd(row.cells[0], LIGHT_BLUE)
        shd(row.cells[1], GREY if i % 2 == 0 else WHITE)
        shd(row.cells[2], GREY if i % 2 == 0 else WHITE)
    doc.add_page_break()

    # ══ 3. TEST SCENARIOS ════════════════════════════════════════════
    h1(doc, "3.  Test Scenarios")
    para(doc, (
        "There are 5 test scenarios. Run them in order. Each scenario takes approximately 5-10 minutes. "
        "For each step: tick Pass if the outcome matches what is described. Tick Fail if it does not."
    ), 10)
    doc.add_paragraph()

    # ── SCENARIO 1: Happy Path ──────────────────────────────────────
    h2(doc, "Scenario 1 — Normal PO: Everything Works Correctly")
    p_box = doc.add_paragraph()
    shd_r = p_box.add_run("🎯  Purpose: Confirm that a normal, complete PO is processed fully automatically — "
                           "Sales Order created in SAP, no manual action needed.")
    shd_r.bold = True; shd_r.font.size = Pt(10)

    doc.add_paragraph()
    para(doc, "📋  Test Setup: Use Test PO 1 (valid PO, known customer, known product code).", 10, bold=True)
    doc.add_paragraph()

    s1 = [
        ("1", "Email Test PO 1 as a PDF attachment to the shared PO mailbox.\n"
              "(Just send a normal email with the PDF attached — exactly as a customer would.)",
         "Email arrives in the shared mailbox inbox.", None),
        ("2", "Ask the AI team to start the processing run.\n"
              "(You do not need to do anything — the system runs automatically.)",
         "System starts processing. You will see progress messages on the AI team's screen.", None),
        ("3", "After processing completes (2–3 minutes), open Outlook and go to the shared mailbox.",
         "✅ The PO email has moved from the Inbox to the 'Processed POs' folder.\n"
         "✅ It is NOT in 'Exception POs' or still in the Inbox.", None),
        ("4", "Check your Customer Service inbox.",
         "✅ You have NOT received any exception or error email for this PO.\n"
         "(No news is good news — exception emails only arrive when something fails.)", None),
        ("5", "Open SAP and search for a new Sales Order for this customer.",
         "✅ A new Sales Order exists in SAP with the correct:\n"
         "   • Customer (Sold-To matches the PO)\n"
         "   • Product / Material\n"
         "   • Quantity and unit\n"
         "   • Requested delivery date", None),
        ("6", "Check the Robona RPA service mailbox.",
         "✅ An automated email has arrived with the subject containing the original PO subject "
         "and the SAP Sales Order number.\n"
         "(This confirms Robona has been notified to link the PO email to the SAP order.)", None),
    ]
    scenario_table(doc, s1, [GREEN if i % 2 == 0 else WHITE for i in range(len(s1))])
    doc.add_paragraph()
    result_row(doc)
    doc.add_page_break()

    # ── SCENARIO 2: Unknown Customer ───────────────────────────────
    h2(doc, "Scenario 2 — Unknown Customer: Exception Alert Sent to CS Team")
    p_box2 = doc.add_paragraph()
    r2b = p_box2.add_run("🎯  Purpose: Confirm that when the system cannot recognise the customer on the PO, "
                          "it does NOT create a Sales Order and instead sends a clear alert to the CS team.")
    r2b.bold = True; r2b.font.size = Pt(10)

    doc.add_paragraph()
    para(doc, "📋  Test Setup: Use Test PO 2 (PO from a customer name NOT in SAP).", 10, bold=True)
    doc.add_paragraph()

    s2 = [
        ("1", "Email Test PO 2 as a PDF attachment to the shared PO mailbox.",
         "Email arrives in the shared mailbox inbox.", None),
        ("2", "Ask the AI team to start the processing run.",
         "System starts processing.", None),
        ("3", "After processing, open Outlook and go to the shared mailbox.",
         "✅ The PO email has moved to 'Exception POs'.\n"
         "❌ It should NOT be in 'Processed POs' and NOT still in the Inbox.", None),
        ("4", "Check the Customer Service team email inbox.",
         "✅ An exception alert email has arrived.\n"
         "The email should clearly show:\n"
         "   • The customer name as extracted from the PO\n"
         "   • A message stating the customer could not be found\n"
         "   • A recommended action for the CS team (e.g. check/add the customer in SAP)\n"
         "   • The original PO PDF attached to the email", None),
        ("5", "Check SAP for a new Sales Order for this customer.",
         "✅ NO Sales Order has been created in SAP.\n"
         "(The system correctly blocked the order — the CS team must handle it manually.)", None),
        ("6", "Check the Robona RPA service mailbox.",
         "✅ NO Robona notification email has been received for this PO.\n"
         "(Robona is only notified on successful SO creation.)", None),
    ]
    scenario_table(doc, s2, [RED if i % 2 == 0 else AMBER for i in range(len(s2))])
    doc.add_paragraph()
    result_row(doc)
    doc.add_page_break()

    # ── SCENARIO 3: Missing Material ───────────────────────────────
    h2(doc, "Scenario 3 — Product Code Not Found: Exception Alert with Missing Items List")
    p_box3 = doc.add_paragraph()
    r3b = p_box3.add_run("🎯  Purpose: Confirm that when the customer is recognised but the product code on the PO "
                          "is not in SAP, the system sends an exception email listing exactly which products are missing.")
    r3b.bold = True; r3b.font.size = Pt(10)

    doc.add_paragraph()
    para(doc, "📋  Test Setup: Use Test PO 3 (known customer, but product code not in SAP).", 10, bold=True)
    doc.add_paragraph()

    s3 = [
        ("1", "Email Test PO 3 as a PDF attachment to the shared PO mailbox.",
         "Email arrives in the shared mailbox inbox.", None),
        ("2", "Ask the AI team to start the processing run.",
         "System starts processing.", None),
        ("3", "After processing, open Outlook.",
         "✅ The PO email has moved to 'Exception POs'.\n"
         "❌ It should NOT be in 'Processed POs'.", None),
        ("4", "Check the Customer Service team email inbox.",
         "✅ An exception alert email has arrived showing:\n"
         "   • Customer: RECOGNISED ✅ (with customer ID shown)\n"
         "   • Product(s): NOT FOUND ❌ — a list of the missing product codes is shown\n"
         "   • Recommended action: add the product mapping to SAP, then resubmit the PO\n"
         "   • The original PO PDF is attached", None),
        ("5", "Check SAP for a Sales Order.",
         "✅ NO Sales Order has been created.\n"
         "(The system did not create a partial order — it blocked the whole PO and flagged to CS.)", None),
        ("6", "Confirm the missing product list in the email.",
         "✅ Every product code that was unrecognised is listed clearly in the email.\n"
         "✅ The product description from the PO is shown alongside each missing code.", None),
    ]
    scenario_table(doc, s3, [RED if i % 2 == 0 else AMBER for i in range(len(s3))])
    doc.add_paragraph()
    result_row(doc)
    doc.add_page_break()

    # ── SCENARIO 4: One email, two POs, one fails ──────────────────
    h2(doc, "Scenario 4 — One Email, Two POs: Partial Failure Handled Correctly")
    p_box4 = doc.add_paragraph()
    r4b = p_box4.add_run("🎯  Purpose: Confirm that when one email contains TWO Purchase Orders and only one "
                          "can be processed, the system creates a Sales Order for the good one AND sends a "
                          "separate exception alert for the failing one.")
    r4b.bold = True; r4b.font.size = Pt(10)

    doc.add_paragraph()
    para(doc, "📋  Test Setup: Send one email with TWO PDF attachments — PO 1 (valid) and PO 3 (missing product code).", 10, bold=True)
    doc.add_paragraph()

    s4 = [
        ("1", "Send ONE email with BOTH Test PO 1 and Test PO 3 attached as PDFs.",
         "Email with two attachments arrives in the shared mailbox.", None),
        ("2", "Ask the AI team to start the processing run.",
         "System processes both PDFs from the same email.", None),
        ("3", "Check Outlook.",
         "✅ The email is in 'Exception POs' (because at least one PO failed).\n"
         "The successfully processed PO 1 has still been pushed to SAP.", None),
        ("4", "Check SAP.",
         "✅ A Sales Order has been created for PO 1 (the valid one).\n"
         "❌ NO Sales Order has been created for PO 3 (the missing product code one).", None),
        ("5", "Check the CS team email inbox.",
         "✅ An exception alert email has arrived specifically for PO 3 (the failing one).\n"
         "✅ PO 1 (the good one) did NOT generate an exception email.", None),
        ("6", "Check the Robona mailbox.",
         "✅ A Robona notification was sent for PO 1 (the successful one).\n"
         "❌ No Robona notification for PO 3 (it did not complete).", None),
    ]
    scenario_table(doc, s4, [AMBER if i % 2 == 0 else WHITE for i in range(len(s4))])
    doc.add_paragraph()
    result_row(doc)
    doc.add_page_break()

    # ── SCENARIO 5: Stage 2 - SAP Block ───────────────────────────
    h2(doc, "Scenario 5 — SAP Block After Submission: CS Team Is Alerted")
    p_box5 = doc.add_paragraph()
    r5b = p_box5.add_run("🎯  Purpose: Confirm that when a Sales Order IS pushed to SAP but SAP rejects or blocks it "
                          "(e.g. credit hold), the CS team receives a second-stage alert email so they can "
                          "resolve the block manually.")
    r5b.bold = True; r5b.font.size = Pt(10)

    doc.add_paragraph()
    para(doc, "📋  Test Setup: Use Test PO 1 (valid PO). Ask the SAP/Celonis team to simulate a credit block for this customer BEFORE the Action Flow runs.", 10, bold=True)
    doc.add_paragraph()

    s5 = [
        ("1", "Process Test PO 1 (the valid PO). Confirm the system has submitted the order to Celonis for SAP processing.",
         "Order data has been pushed to the system. No exception email at this stage.", None),
        ("2", "Ask the SAP/Celonis team to set a credit block or validation error for this customer so SAP rejects the order.",
         "Block is in place in SAP before the Action Flow runs.", None),
        ("3", "Ask the AI team to run the post-processing step (this checks Celonis for results).",
         "System detects the block / failure from the Celonis Action Flow result.", None),
        ("4", "Check the CS team email inbox.",
         "✅ A second exception alert email has arrived showing:\n"
         "   • Customer name and PO number\n"
         "   • Reason for failure (e.g. credit block)\n"
         "   • Recommended action (e.g. release the credit block in SAP, then re-process)", None),
        ("5", "Check Outlook.",
         "✅ The original PO email has been moved to 'Exception POs'.", None),
        ("6", "Check the Robona mailbox.",
         "✅ NO Robona notification has been sent.\n"
         "(Robona only receives a notification when the Sales Order is fully created with NO blocks.)", None),
    ]
    scenario_table(doc, s5, [RED if i % 2 == 0 else AMBER for i in range(len(s5))])
    doc.add_paragraph()
    result_row(doc)
    doc.add_page_break()

    # ══ 4. OVERALL SUMMARY ═══════════════════════════════════════════
    h1(doc, "4.  UAT Overall Summary")
    para(doc, "Complete this table after all 5 scenarios have been run.", 10)
    doc.add_paragraph()

    sum_tbl = doc.add_table(rows=7, cols=4)
    sum_tbl.style = "Table Grid"
    hdr(sum_tbl, ["Scenario", "Description", "Result", "Client Sign-Off"], DARK_BLUE, WHITE, 9)
    col_widths(sum_tbl, [2.0, 7.5, 2.5, 4.0])
    scenarios = [
        ("Scenario 1", "Normal PO — Full Sales Order created in SAP"),
        ("Scenario 2", "Unknown Customer — Exception email sent, no SO created"),
        ("Scenario 3", "Missing Product Code — Exception with missing items list"),
        ("Scenario 4", "Two POs in one email — One succeeds, one exception raised"),
        ("Scenario 5", "SAP Block — Second-stage alert sent to CS team"),
    ]
    bgs = [GREEN, RED, RED, AMBER, RED]
    for i, (scen, desc) in enumerate(scenarios):
        row = sum_tbl.rows[i + 1]
        cell_text(row.cells[0], scen, 9, bold=True)
        cell_text(row.cells[1], desc, 9)
        cell_text(row.cells[2], "☐  PASS\n☐  FAIL", 9)
        cell_text(row.cells[3], "☐  ACCEPTED\n☐  REJECTED\n\nName: ___________\nDate:  ___________", 9)
        for c in row.cells:
            shd(c, bgs[i])
    doc.add_paragraph()

    para(doc, "UAT is ACCEPTED when:", 10, bold=True)
    bullet(doc, "All 5 scenarios are marked PASS by the tester.")
    bullet(doc, "All 5 scenarios are marked ACCEPTED by the client.")
    bullet(doc, "All gaps recorded in the Gap Log below have an agreed resolution date.")
    doc.add_page_break()

    # ══ 5. SIGN-OFF ══════════════════════════════════════════════════
    h1(doc, "5.  Official Sign-Off")
    para(doc, "By signing below, both parties confirm that UAT has been completed and the results are accepted.", 10)
    doc.add_paragraph()

    so_tbl = doc.add_table(rows=4, cols=4)
    so_tbl.style = "Table Grid"
    hdr(so_tbl, ["Role", "Full Name", "Signature", "Date"], DARK_BLUE, WHITE, 9)
    col_widths(so_tbl, [5.0, 4.0, 4.0, 3.0])
    roles = ["Client — Project Lead", "Client — CS Team Lead", "AI Project Lead (Supplier)"]
    for i, role in enumerate(roles):
        row = so_tbl.rows[i + 1]
        cell_text(row.cells[0], role, 9, bold=True)
        shd(row.cells[0], LIGHT_BLUE)
        for c in row.cells[1:]:
            shd(c, GREY if i % 2 == 0 else WHITE)
    doc.add_page_break()

    # ══ 6. GAP LOG ═══════════════════════════════════════════════════
    h1(doc, "6.  Gap & Action Log")
    para(doc, (
        "Use this table to record anything that did not pass during UAT. "
        "Each gap must have an owner and a target resolution date before UAT can be closed."
    ), 10)
    doc.add_paragraph()

    gap_tbl = doc.add_table(rows=7, cols=5)
    gap_tbl.style = "Table Grid"
    hdr(gap_tbl, ["Ref", "What Happened (Actual)", "What Was Expected", "Who Will Fix It", "Target Date"], DARK_BLUE, WHITE, 9)
    col_widths(gap_tbl, [1.0, 4.5, 4.5, 3.0, 3.0])
    for i in range(1, 7):
        row = gap_tbl.rows[i]
        cell_text(row.cells[0], f"G-{i:02d}", 9, bold=True)
        shd(row.cells[0], LIGHT_BLUE)
        for c in row.cells[1:]:
            shd(c, GREY if i % 2 == 0 else WHITE)

    out = "Touchless_Order_Creation_UAT_Functional.docx"
    doc.save(out)
    print(f"Saved: {out}")

if __name__ == "__main__":
    build()
