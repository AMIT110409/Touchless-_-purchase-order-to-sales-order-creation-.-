import os
import sys
import json
import sqlite3
import pandas as pd
from pycelonis import get_celonis
import pycelonis.pql as pql
from pycelonis.pql.saola_connector import AnalysisSaolaConnector
from saolapy.pql.base import PQLFilter
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def main():
    print("=" * 70)
    print("  CELONIS CUSTOMER GROUP & SALESPERSON EMAIL EXTRACTION")
    print("=" * 70)

    # ---------------------------------------------------------
    # STEP 1: Gather active customer IDs
    # ---------------------------------------------------------
    print("\n[1/5] Gathering active customer IDs...")
    customers = set()

    # 1.1 From EXPORT Excel file
    excel_path = "EXPORT_20260216_100124.XLSX"
    if os.path.exists(excel_path):
        try:
            df_exp = pd.read_excel(excel_path)
            if 'Customer' in df_exp.columns:
                for val in df_exp['Customer'].dropna().unique():
                    val_str = str(int(val) if isinstance(val, (int, float)) else val).strip()
                    if val_str.isdigit() and len(val_str) == 10 and val_str.startswith("402"):
                        customers.add(val_str)
            print(f"  - Found {len(customers)} customer IDs in {excel_path}")
        except Exception as e:
            print(f"  - Error reading {excel_path}: {e}")
    else:
        print(f"  - Warning: {excel_path} not found")

    # 1.2 From all jsonl log files (skipping huge virtual env folders)
    jsonl_count = 0
    skip_dirs = {'venv', 'venv311', '.git', '__pycache__', '.ipynb_checkpoints', 'local_test_folder', 'unprocessed_run', 'Stebro_reextract'}
    
    for root, dirs, files in os.walk("."):
        # Modify dirs in-place to prevent os.walk from entering skipped directories
        dirs[:] = [d for d in dirs if d not in skip_dirs]
        
        for file in files:
            if file.endswith(".jsonl"):
                path = os.path.join(root, file)
                try:
                    with open(path, 'r', encoding='utf-8') as f:
                        for line in f:
                            if not line.strip():
                                continue
                            try:
                                data = json.loads(line)
                                header = data.get('header_fields', {})
                                for k, v in header.items():
                                    if 'customer' in k.lower() or 'sold_to' in k.lower():
                                        val = str(v).strip()
                                        if val.isdigit() and len(val) == 10 and val.startswith("402"):
                                            customers.add(val)
                                val = data.get('matched_customer_id')
                                if val:
                                    val_str = str(val).strip()
                                    if val_str.isdigit() and len(val_str) == 10 and val_str.startswith("402"):
                                        customers.add(val_str)
                            except Exception:
                                pass
                    jsonl_count += 1
                except Exception:
                    pass
    print(f"  - Scanned jsonl files (skipped huge directories): found total of {len(customers)} unique customer IDs")

    cust_list = sorted(list(customers))
    print(f"  - Unified Customer IDs list ({len(cust_list)} items):")
    print(f"    {cust_list}")

    if not cust_list:
        print("ERROR: No customer IDs found to process!")
        return

    # ---------------------------------------------------------
    # STEP 2: Connect to Celonis and run PQL query
    # ---------------------------------------------------------
    print("\n[2/5] Connecting to Celonis Cloud...")
    url = 'https://envalior-sb.eu-1.celonis.cloud/'
    api_token = 'NTFkY2QyZTQtMDQ4OS00MTljLThhMGUtMWJkMGRlOTUxNTkzOk5KM0RITkh2eVJvZmt6Z0ZvU05tT0s2MHRwMnc5cGhtaGw1a1l3L0tkTFAx'

    try:
        c = get_celonis(url, api_token)
        print("  - Connection established successfully.")
        
        print("  - Locating OCPM Space and Touchless Order Creation assets...")
        space = c.apps.get_spaces().find('OCPM')
        data_pool = c.data_integration.get_data_pool('663e7e2d-74f0-4cf1-a0b4-ca4faf5fee06')
        package = space.get_packages().find('Touchless Order Creation')
        data_model = data_pool.get_data_model('3f193f92-a398-4989-915c-cc100e67d421')
        analysis = package.get_analysis('61ec7315-0e7a-467a-8b89-116d78e1ca9e#!')
        content = analysis.get_content()
        sheets = content.draft.document.components

        print("  - Extracting query for component '5ba818be-4fda-46a1-96ee-1d1fd9060d60' (New Sheet 2)...")
        sheet = sheets.find('New Sheet 2')
        comp = sheet.components.find_by_id('5ba818be-4fda-46a1-96ee-1d1fd9060d60')
        query = comp.get_query()

        # Add customer list filter
        cust_str = ", ".join([f"'{cust}'" for cust in cust_list])
        filter_expr = f'FILTER "o_celonis_Customer"."Number" IN ({cust_str})'
        print(f"  - Appending filter: {filter_expr}")
        query.filters.append(PQLFilter(query=filter_expr))

        print("  - Executing query against Data Model...")
        df_raw = pql.DataFrame.from_pql(query, saola_connector=AnalysisSaolaConnector(data_model, analysis)).to_pandas()
        print(f"  - Successfully fetched {len(df_raw)} raw records from Celonis.")

    except Exception as e:
        print(f"ERROR connecting or querying Celonis: {e}")
        return

    # ---------------------------------------------------------
    # STEP 3: Normalize and Clean Data
    # ---------------------------------------------------------
    print("\n[3/5] Normalizing extracted data...")
    
    # Map raw technical column names to reader friendly column names
    col_mapping = {
        'Number': 'Customer ID',
        'Name': 'Customer Name',
        'Sold to': 'Sold To',
        '#{o_custom_CustomerRoleMaterial.Ship to}': 'Ship To',
        '#{o_custom_CustomerRoleMaterial.Ship to name}': 'Ship To Name',
        '#{o_custom_CustomerRoleMaterial.Material}': 'Material (Internal)',
        'MaterialDescription': 'Material Description',
        'CustomerMaterialNumber': 'Customer Material Number',
        '#{o_custom_CustomerRoleMaterial.CustomerMaterialDescription}': 'Customer Material Description',
        'SalesOrganization': 'Sales Organization',
        '#{o_custom_CustomerRoleMaterial.CustomerGroup2}': 'Customer Group 2',
        '#{o_custom_CustomerRoleMaterial.SalespersonEmail}': 'Salesperson Email'
    }

    # Ensure all mapped columns exist in retrieved data frame
    cols_to_use = [col for col in col_mapping.keys() if col in df_raw.columns]
    df_clean = df_raw[cols_to_use].rename(columns={k: v for k, v in col_mapping.items() if k in cols_to_use})

    # Normalize Sold To and Ship To
    if 'Sold To' in df_clean.columns:
        df_clean['Sold To'] = df_clean['Sold To'].fillna('-').astype(str).str.strip()
        df_clean['Sold To'] = df_clean['Sold To'].apply(lambda x: '-' if str(x).strip() in ('', 'None', 'nan', '<NA>') else str(x).strip())
    else:
        df_clean['Sold To'] = '-'

    if 'Ship To' in df_clean.columns:
        df_clean['Ship To'] = df_clean['Ship To'].fillna('-').astype(str).str.strip()
        df_clean['Ship To'] = df_clean['Ship To'].apply(lambda x: '-' if str(x).strip() in ('', 'None', 'nan', '<NA>') else str(x).strip())
    else:
        df_clean['Ship To'] = '-'

    # Normalize Customer Group 2: Empty/nan/None should be '-' to match Celonis UI
    if 'Customer Group 2' in df_clean.columns:
        df_clean['Customer Group 2'] = df_clean['Customer Group 2'].fillna('-')
        df_clean['Customer Group 2'] = df_clean['Customer Group 2'].apply(lambda x: '-' if str(x).strip() in ('', 'None', 'nan', '<NA>') else str(x).strip())
    else:
        df_clean['Customer Group 2'] = '-'

    # Strip whitespace from salesperson email addresses
    if 'Salesperson Email' in df_clean.columns:
        df_clean['Salesperson Email'] = df_clean['Salesperson Email'].fillna('-').astype(str).str.strip()
    else:
        df_clean['Salesperson Email'] = '-'

    # Create Summary DataFrame (Include Sold To and Ship To)
    summary_cols = ['Customer ID', 'Customer Name', 'Sold To', 'Ship To', 'Customer Group 2', 'Salesperson Email']
    df_summary = df_clean[summary_cols].drop_duplicates().copy()
    
    # Group by Customer ID and Name and combine different values if they exist (e.g. some materials have Z01, some -)
    def aggregate_unique(series):
        unique_vals = sorted(list(set(str(val) for val in series.dropna() if str(val).strip() != '')))
        if not unique_vals:
            return '-'
        return ", ".join(unique_vals)

    df_summary_agg = df_summary.groupby(['Customer ID', 'Customer Name']).agg({
        'Sold To': aggregate_unique,
        'Ship To': aggregate_unique,
        'Customer Group 2': aggregate_unique,
        'Salesperson Email': aggregate_unique
    }).reset_index()

    # Sort
    df_summary_agg.sort_values(by='Customer ID', inplace=True)
    df_clean.sort_values(by=['Customer ID', 'Material (Internal)'], inplace=True)

    print(f"  - Generated aggregate summary for {len(df_summary_agg)} unique customers.")
    print(f"  - Detailed table contains {len(df_clean)} rows.")

    # ---------------------------------------------------------
    # STEP 4: Write to styled Excel workbook
    # ---------------------------------------------------------
    output_excel = "extracted_customer_groups_all.xlsx"
    print(f"\n[4/5] Exporting styled Excel file to '{output_excel}'...")

    try:
        writer = pd.ExcelWriter(output_excel, engine='openpyxl')
        
        # Write to sheets
        df_summary_agg.to_excel(writer, sheet_name='Customer Group Summary', index=False)
        df_clean.to_excel(writer, sheet_name='Detailed Table', index=False)
        
        # Access openpyxl objects for styling
        workbook = writer.book
        
        # Define Premium Styles
        header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        data_font = Font(name='Segoe UI', size=10, color='333333')
        title_font = Font(name='Segoe UI', size=16, bold=True, color='1B365D')
        
        header_fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid') # Sleek navy
        zebra_fill = PatternFill(start_color='F5F7FA', end_color='F5F7FA', fill_type='solid')  # Very light grey-blue
        accent_fill = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
        
        thin_border_side = Side(border_style="thin", color="CCCCCC")
        thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        
        align_left = Alignment(horizontal='left', vertical='center')
        align_center = Alignment(horizontal='center', vertical='center')
        
        for name in workbook.sheetnames:
            ws = workbook[name]
            
            # Show grid lines explicitly
            ws.views.sheetView[0].showGridLines = True
            
            # Format row heights
            ws.row_dimensions[1].height = 28 # Header row
            
            # Style header row
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = align_center if 'ID' in str(cell.value) or 'Group' in str(cell.value) or 'Organization' in str(cell.value) else align_left
                cell.border = thin_border
            
            # Style data rows (Only style cells individually for the summary sheet, detailed is too large)
            if name == 'Customer Group Summary':
                for row in range(2, ws.max_row + 1):
                    ws.row_dimensions[row].height = 20
                    is_zebra = (row % 2 == 0)
                    
                    for col in range(1, ws.max_column + 1):
                        cell = ws.cell(row=row, column=col)
                        cell.font = data_font
                        cell.border = thin_border
                        
                        # Zebra striping
                        if is_zebra:
                            cell.fill = zebra_fill
                        
                        # Alignment
                        col_name = str(ws.cell(row=1, column=col).value)
                        if col_name in ('Customer ID', 'Sold To', 'Ship To', 'Sales Organization', 'Customer Group 2'):
                            cell.alignment = align_center
                        else:
                            cell.alignment = align_left
                            
                        # Highlight Z01 cells in light green/teal for premium visibility
                        if col_name == 'Customer Group 2' and str(cell.value).strip() == 'Z01':
                            cell.fill = PatternFill(start_color='E6FFFA', end_color='E6FFFA', fill_type='solid')
                            cell.font = Font(name='Segoe UI', size=10, bold=True, color='0D9488')
            
            # Adjust column widths automatically
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Save workbook
        writer.close()
        print(f"  - Successfully created formatted Excel workbook: {output_excel}")

    except Exception as e:
        print(f"ERROR exporting Excel workbook: {e}")
        return

    # ---------------------------------------------------------
    # STEP 5: Verification & Summary Report
    # ---------------------------------------------------------
    print("\n[5/5] Performing verification check...")
    if os.path.exists(output_excel):
        print(f"\n{'-'*60}")
        print(f"  VERIFICATION SUCCESSFUL: file exists.")
        print(f"  Path: {os.path.abspath(output_excel)}")
        print(f"{'-'*60}\n")
        
        # Display short summary table on screen
        print(f"{'Customer ID':<12} | {'Customer Name':<35} | {'Sold To':<12} | {'Ship To':<20} | {'Group 2':<8} | {'Salesperson Emails'}")
        print(f"{'-'*120}")
        for idx, row in df_summary_agg.iterrows():
            cust_id = str(row['Customer ID'])
            cust_name = str(row['Customer Name'])
            if len(cust_name) > 32:
                cust_name = cust_name[:29] + "..."
            sold_to = str(row['Sold To'])
            if len(sold_to) > 10:
                sold_to = sold_to[:8] + "..."
            ship_to = str(row['Ship To'])
            if len(ship_to) > 18:
                ship_to = ship_to[:15] + "..."
            group = str(row['Customer Group 2'])
            emails = str(row['Salesperson Email'])
            if len(emails) > 30:
                emails = emails[:27] + "..."
            
            # Print with ignore encoding errors to prevent console crashes
            line = f"{cust_id:<12} | {cust_name:<35} | {sold_to:<12} | {ship_to:<20} | {group:<8} | {emails}"
            print(line.encode(sys.stdout.encoding, errors='replace').decode(sys.stdout.encoding))
        print(f"{'-'*120}\n")
    else:
        print("ERROR: Excel file could not be found after generation.")

if __name__ == "__main__":
    main()
