from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_unit_1():
    doc = Document()
    
    # Title
    title = doc.add_heading('Unit 1: Executive Overview', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph(
        "The Envalior AI-Driven Purchase Order (PO) Workflow is a premium, touchless automation solution "
        "designed to eliminate manual data entry in the sales order creation process. By leveraging "
        "State-of-the-Art AI (Claude Sonnet), GPU-accelerated OCR (PaddleOCR), and enterprise-grade "
        "validation (Celonis), the system ensures 100% traceability and significant reduction in order processing time."
    )
    
    doc.add_heading('2. High-Level Process Flow', level=1)
    doc.add_paragraph("[IMAGE PLACEHOLDER: BUSINESS_ARCHITECTURE_DIAGRAM]")
    doc.add_paragraph(
        "The workflow follows a 7-phase architecture:\n"
        "1. TRIGGER: Automated email polling via MS Graph API.\n"
        "2. OCR & EXTRACTION: Intelligent data capture from any PO format.\n"
        "3. KB ENRICHMENT: Customer alias resolution using a 52K+ record Knowledge Base.\n"
        "4. MATERIAL MAPPING: Automated cross-referencing of internal vs. external material codes.\n"
        "5. VALIDATION: Real-time business rule checks via the Celonis Action Flow.\n"
        "6. OUTPUT & ACTION: Seamless Sales Order creation in SAP EPP.\n"
        "7. EXCEPTIONS: Smart routing to regional CS teams for low-confidence files."
    )
    
    doc.add_heading('3. Key Business Objectives', level=1)
    objectives = [
        "Touchless Automation: Transform raw PDFs into SAP orders without human intervention.",
        "Total Traceability: Original emails are archived and attached directly to SAP SO records via Robona RPA.",
        "Regional Efficiency: Automated exception routing to EMEA, APAC, and AMS desks.",
        "Feedback Loop: Human-in-the-loop corrections continuously retrain the AI Knowledge Base."
    ]
    for obj in objectives:
        doc.add_paragraph(obj, style='List Bullet')

    doc.save('Unit_1_Executive_Overview.docx')
    print("Unit 1 created.")

if __name__ == "__main__":
    create_unit_1()
